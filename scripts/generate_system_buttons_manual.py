"""
Generate a comprehensive Microsoft Word (.docx) document detailing every button,
form submission, interactive control, workflow, and system output for every user role
in the SafeWaste Medical Waste Management Platform.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set the background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table, color="CBD5E1"):
    """Set clean subtle borders for a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def make_callout(doc, title, text, tone="info"):
    """Create a styled callout box."""
    tones = {
        "info": {"border": "2563EB", "bg": "EFF6FF", "title_color": RGBColor(30, 64, 175)},
        "success": {"border": "059669", "bg": "ECFDF5", "title_color": RGBColor(6, 95, 70)},
        "warning": {"border": "D97706", "bg": "FFFBEB", "title_color": RGBColor(146, 64, 14)},
        "danger": {"border": "DC2626", "bg": "FEF2F2", "title_color": RGBColor(153, 27, 27)},
    }
    cfg = tones.get(tone, tones["info"])

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)

    # Thick left border
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg["border"]}"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    set_cell_background(cell, cfg["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"• {title}")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = cfg["title_color"]

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    run_body = p2.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_button_table(doc, button_data, table_title=None):
    """
    Render a standard 4-column Button & Workflow specification table:
    Col 0: Button & Visual Element (1.5 in)
    Col 1: Purpose & Trigger Condition (1.4 in)
    Col 2: Technical Workflow & Architecture (2.2 in)
    Col 3: System Output & Result (1.9 in)
    """
    if table_title:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(8)
        p_sub.paragraph_format.space_after = Pt(3)
        run_sub = p_sub.add_run(table_title)
        run_sub.bold = True
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(11)
        run_sub.font.color.rgb = RGBColor(15, 118, 110)

    headers = [
        "Button / Control Element",
        "Purpose & Conditions",
        "Triggered Workflow & Logic",
        "System Output & Downstream Effects"
    ]
    col_widths = [Inches(1.5), Inches(1.4), Inches(2.2), Inches(1.9)]

    table = doc.add_table(rows=len(button_data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "CBD5E1")

    # Format header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Format data rows
    for row_idx, data in enumerate(button_data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"

        for col_idx in range(4):
            cell = row_cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1

            # First column has bold title
            text_val = data[col_idx]
            lines = text_val.split("\n")
            for line_idx, line in enumerate(lines):
                if line_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.1

                if col_idx == 0 and line_idx == 0:
                    r = p.add_run(line)
                    r.bold = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(15, 23, 42)
                elif line.startswith("•") or line.startswith("-") or line.startswith("API:") or line.startswith("Action:"):
                    r = p.add_run(line)
                    r.font.name = "Calibri"
                    r.font.size = Pt(8.5)
                    if line.startswith("API:") or line.startswith("Action:"):
                        r.bold = True
                        r.font.color.rgb = RGBColor(14, 116, 144)
                    else:
                        r.font.color.rgb = RGBColor(51, 65, 85)
                else:
                    r = p.add_run(line)
                    r.font.name = "Calibri"
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = Document()

    # Page setup: Standard Letter, 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # DOCUMENT COVER / TITLE HEADER
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_sub = title_p.add_run("SAFEWASTE HEALTHCARE LOGISTICS PLATFORM · TECHNICAL & USER MANUAL")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.bold = True
    run_sub.font.color.rgb = RGBColor(15, 118, 110) # Teal

    h1_p = doc.add_paragraph()
    h1_p.paragraph_format.space_before = Pt(4)
    h1_p.paragraph_format.space_after = Pt(6)
    run_h1 = h1_p.add_run("Complete Interactive Buttons, Operational Workflows & System Output Manual")
    run_h1.font.name = "Calibri"
    run_h1.font.size = Pt(22)
    run_h1.bold = True
    run_h1.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_before = Pt(0)
    desc_p.paragraph_format.space_after = Pt(12)
    desc_p.paragraph_format.line_spacing = 1.15
    run_desc = desc_p.add_run(
        "A comprehensive reference specification detailing every button, form trigger, modal action, "
        "and interactive control across all user personas (Super Administrator, Hospital Administrator, "
        "Hospital Staff / EHO, Fleet Driver, Disposal Facility Operator, EFDA Regulatory Officer, "
        "and Public Verification Visitors). Includes technical execution pipelines, database transitions, "
        "idempotency mechanisms, and downstream operational outputs."
    )
    run_desc.font.size = Pt(10.5)

    # Metadata bar
    meta_tbl = doc.add_table(rows=1, cols=4)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_headers = [
        ("Platform Version", "v1.2 (Oct 2026 Baseline)"),
        ("Active Deployment", "Dire Dawa Regional Hub"),
        ("Audience", "Admins, Operators, Regulators"),
        ("Currency Standard", "ETB (Ethiopian Birr)")
    ]
    for idx, (label, val) in enumerate(meta_headers):
        c = meta_tbl.cell(0, idx)
        c.width = Inches(1.75)
        set_cell_background(c, "F1F5F9")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(f"{label}\n")
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor(100, 116, 139)
        r2 = p.add_run(val)
        r2.font.size = Pt(9)
        r2.bold = True
        r2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 1: SYSTEM ARCHITECTURE & ROLE MATRIX
    # -------------------------------------------------------------
    s1 = doc.add_heading(level=1)
    s1.paragraph_format.space_before = Pt(16)
    s1.paragraph_format.space_after = Pt(6)
    r_s1 = s1.add_run("1. Architectural Foundation & Consolidated Role Matrix")
    r_s1.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "SafeWaste enforces a zero-trust, role-based access control (RBAC) architecture with exactly one active role per user. "
        "Every state-changing operational button is bound to an idempotent execution pipeline, audited in an append-only "
        "cryptographic hash-chain ledger, and coupled to an asynchronous outbox messaging dispatcher. All billing figures and "
        "contracts are strictly denominated in Ethiopian Birr (ETB)."
    )

    role_matrix_data = [
        [
            "Super Administrator\n(super_admin)",
            "Nebiyu (Platform Admin)\nRoute: /admin",
            "• Live GPS fleet radar & tracker management\n• Facility registry, readiness & licensing\n• Trip planning, stop sequencing & driver dispatch\n• Pickup override with audited rationale\n• Monthly invoice generation (ETB) & offline settlement\n• Certificate revocation & global system settings",
            "Full global read/write oversight across all 18 regional healthcare facilities and disposal yard."
        ],
        [
            "Hospital Administrator\n(hospital_admin)",
            "Dr. Aster Tadesse\nRoute: /hospital",
            "• Facility waste ledger & container stage tracking\n• Official destruction certificates feed & download\n• Facility staff management & pickup OTP delegation\n• Medicine stock oversight & disposal authorization\n• Facility monthly invoice summary in ETB",
            "Full read/write scoped strictly to own facility (Dil Chora Referral Hospital). Cannot see other facilities."
        ],
        [
            "Hospital Staff / EHO\n(hospital_staff)",
            "Sr. Genet Alemu\nRoute: /hospital",
            "• Biohazard waste registration by MoH category\n• Barcoded container labeling & label printing\n• Medicine stock uploading & disposal requests\n• Real-time truck-at-door handover monitoring\n• In-app handover confirmation fallback",
            "Operational waste entry & pickup handover scoped strictly to own facility (Sabian General Hospital)."
        ],
        [
            "Fleet Driver\n(driver)",
            "Abebe Kebede\nRoute: /driver",
            "• Mobile-first console optimized for touch\n• Daily assigned trip initiation & stop navigation\n• 'I have arrived' trigger with GPS geofence test\n• Camera QR barcode container scanning\n• SMS OTP verification & pickup completion\n• Exception handling (Skip stop with reason)",
            "Restricted strictly to own assigned vehicle and trips in active progress. Zero access to billing or other drivers."
        ],
        [
            "Disposal Facility Operator\n(disposal_operator)",
            "Eng. Kassahun Worku\nRoute: /disposal",
            "• Central incinerator yard reception desk\n• Inbound collection truck container check-in\n• Missing / Refused container discrepancy logging\n• Incinerator / Autoclave destruction queue\n• Irreversible destruction confirmation & record lock\n• Automatic certificate PDF generation trigger",
            "Terminal execution scoped strictly to the regional disposal and treatment plant. Zero route modification rights."
        ],
        [
            "EFDA Regulatory Officer\n(efda_officer)",
            "Tariku Haile\nRoute: /regulator",
            "• Regional compliance oversight across all 18 facilities\n• Volume analytics & hospital generation league table\n• Historical route replay with tracker telemetry\n• Cryptographic audit ledger verification (SHA-256)\n• Regulatory medicine disposal approval / rejection",
            "Strictly read-only across all data surfaces, with exactly ONE write permission: regulatory approval of expired drug destruction."
        ],
        [
            "Public Visitor\n(unauthenticated)",
            "General Public / External Audit\nRoute: /verify/[token]",
            "• Public certificate authenticity check\n• Container count, weight & destruction method review\n• Official PDF Certificate download\n• Cryptographic SHA-256 PDF file tamper check",
            "Zero sign-in required. Shows anonymized certificate verification data without person names or phone numbers."
        ],
    ]

    add_button_table(doc, role_matrix_data, "Table 1.1: Consolidated Persona & Responsibility Matrix")

    make_callout(
        doc,
        "Binding Architecture Rule: Mandatory Request Pipeline",
        "Every state-changing button documented below follows an identical 6-step backend pipeline: "
        "(1) Clerk session authentication, (2) Database role & facility scope verification, "
        "(3) Zod input validation, (4) Idempotency-Key collision check in PostgreSQL, "
        "(5) Atomic transactional state transition with audit log hash insertion, and "
        "(6) Outbox message dispatch. Double-clicks return cached HTTP responses without re-executing logic.",
        "info"
    )

    # -------------------------------------------------------------
    # SECTION 2: SUPER ADMINISTRATOR (/admin)
    # -------------------------------------------------------------
    s2 = doc.add_heading(level=1)
    s2.paragraph_format.space_before = Pt(16)
    s2.paragraph_format.space_after = Pt(6)
    r_s2 = s2.add_run("2. Super Administrator Persona (/admin) — Buttons & Workflows")
    r_s2.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "The Super Administrator oversees the entire regional medical waste ecosystem. Buttons in this area "
        "control fleet operations, facility onboarding, route scheduling, financial billing in ETB, and emergency overrides."
    )

    # 2.1 Live Fleet Radar & Dashboard
    admin_radar_buttons = [
        [
            "[Refresh Radar]\nHeader Control\nIcon: RefreshCw",
            "Purpose: Forces immediate polling update of active vehicle GPS positions.\n\nCondition: Active trip IN_PROGRESS.",
            "Action: Queries /api/tracking/live for vehicles with active trips.\n• Validates tracker timestamps against tracker_max_clock_skew_minutes.\n• Computes distance to destination facility.",
            "Output: Updates Leaflet map markers, vehicle speed indicators, and breadcrumb trails without page reload."
        ],
        [
            "[Center Vehicle]\nMap Marker\nIcon: Navigation",
            "Purpose: Focuses map camera on a specific truck in transit.\n\nCondition: Vehicle marker visible.",
            "Action: Client-side Leaflet flyTo() centering on vehicle coordinates.\n• Zooms to street level (16x).\n• Expands popup with driver name, plate number, and payload weight.",
            "Output: Smooth map pan and zoom to truck location."
        ],
        [
            "[View Alert Context]\nAlerts Table\nIcon: ExternalLink",
            "Purpose: Directly navigates from a critical system alert to the affected entity.\n\nCondition: Alert present.",
            "Action: Router navigation to targeted view (/admin/trips/[id], /admin/facilities/[code], or /admin/finance).\n• Passes query parameters highlighting the discrepancy.",
            "Output: Instant redirection to the administrative screen where corrective action can be taken."
        ],
    ]
    add_button_table(doc, admin_radar_buttons, "2.1 Fleet Radar & Dashboard Monitoring Controls")

    # 2.2 Facility Management
    admin_facility_buttons = [
        [
            "[Create Facility]\n/admin/facilities/new\nButton: Primary Blue",
            "Purpose: Registers a new health facility or disposal site into master database.\n\nCondition: Form fields valid.",
            "Action: Calls createFacilityAction(values).\n• Validates unique facility code format (XX-XXX-000).\n• Validates lat/lng geocodes.\n• Inserts record into facilities table with status ACTIVE.\n• Logs FACILITY_CREATED in immutable audit ledger.",
            "Output: Redirects to /admin/facilities/[code]. New facility is immediately eligible for contract assignment and user onboarding."
        ],
        [
            "[Save Changes]\n/admin/facilities/[code]\nButton: Primary Blue",
            "Purpose: Updates facility metadata (name, woreda, kebele, contact address, geocoordinates).\n\nCondition: Facility exists.",
            "Action: Calls updateFacilityAction(code, values).\n• Code is immutable and disabled.\n• Updates facility record.\n• Recomputes arrival geofence radius for ongoing trips.\n• Writes audit log entry.",
            "Output: Displays green 'Saved' confirmation banner. Updated coordinates take effect on next driver arrival."
        ],
        [
            "[Change Status]\nFacility Status Control\nButton: Danger Red",
            "Purpose: Transitions facility between ACTIVE, INACTIVE, and SUSPENDED.\n\nCondition: Admin selects status differing from current.",
            "Action: Calls setFacilityStatusAction(code, { status }).\n• If SUSPENDED, cancels open scheduled pickups on PLANNED trips.\n• Blocks new waste registration via database constraint.\n• Writes audit log with administrator ID.",
            "Output: StatusBadge updates instantly. If suspended, alerts appear on hospital dashboard and facility is excluded from trip planning."
        ],
    ]
    add_button_table(doc, admin_facility_buttons, "2.2 Health Facility Administration Controls")

    # 2.3 Fleet, Vehicles, Trackers & Licences
    admin_fleet_buttons = [
        [
            "[Add Vehicle]\n/admin/fleet\nButton: Primary Blue",
            "Purpose: Enrolls a new collection truck into the active fleet.\n\nCondition: Unique plate number.",
            "Action: Calls createVehicleAction(values).\n• Validates plate format and payload capacity (kg).\n• Inserts vehicle row with status ACTIVE.\n• Audit log: VEHICLE_CREATED.",
            "Output: Vehicle appears in fleet list and becomes available in Trip Planner selector."
        ],
        [
            "[Change Status]\nVehicle Table Row\nDropdown: Select",
            "Purpose: Sets vehicle state to ACTIVE, MAINTENANCE, or OFFLINE.\n\nCondition: Disabled if IN_TRANSIT.",
            "Action: Calls setVehicleStatusAction(vehicleId, { status }).\n• Prevents changes if vehicle is currently on an active trip.\n• If MAINTENANCE, excludes vehicle from trip planning dropdown.",
            "Output: Vehicle status badge updates; scheduled alerts clear if returned to ACTIVE."
        ],
        [
            "[Register Tracker]\n/admin/fleet\nButton: Primary Blue",
            "Purpose: Associates a hardware GPS tracking device with a vehicle.\n\nCondition: Valid hardware device ID.",
            "Action: Calls createTrackerAction(values).\n• Generates high-entropy cryptographic ingest token.\n• Stores SHA-256 hash in database.\n• Binds deviceId to selected vehicle plate.",
            "Output: Displays modal [TokenOnce] with raw ingest token. Token is shown exactly once and cannot be retrieved again."
        ],
        [
            "[New Token]\nTracker Table Row\nButton: Secondary",
            "Purpose: Rotates device token if compromised or replaced.\n\nCondition: Tracker registered.",
            "Action: Calls rotateTrackerTokenAction(trackerId).\n• Invalidates previous token hash.\n• Generates new device key.\n• Writes security audit entry.",
            "Output: Opens [TokenOnce] dialog. Previous token fails ingest with HTTP 401 immediately."
        ],
        [
            "[Deactivate / Activate]\nTracker Table Row\nButton: Secondary",
            "Purpose: Toggles tracker status.\n\nCondition: Tracker exists.",
            "Action: Calls setTrackerActiveAction(trackerId, !active).\n• If deactivated, ingest endpoint rejects coordinates from this device.",
            "Output: StatusBadge switches between ACTIVE / INACTIVE; silent tracker alerts suppressed if inactive."
        ],
        [
            "[Add Licence]\n/admin/fleet\nButton: Primary Blue",
            "Purpose: Records transport permit, driver licence, or facility operational permit.\n\nCondition: Valid expiry date.",
            "Action: Calls addLicenceAction(values).\n• Categorizes under VEHICLE, DRIVER, FACILITY, or SITE.\n• Sets expiration warning threshold (default 30 days).\n• Inserts licence record.",
            "Output: Licence row added. Status calculated automatically (VALID, EXPIRING, EXPIRED). If expired, trip planner blocks assignment."
        ],
    ]
    add_button_table(doc, admin_fleet_buttons, "2.3 Fleet, Hardware Trackers & Compliance Licencing Controls")

    # 2.4 User Access & OTP Recipient Delegation
    admin_user_buttons = [
        [
            "[Save Access]\n/admin/users/[id]\nButton: Primary Blue",
            "Purpose: Modifies user role, facility assignment, active login status, and pickup OTP authorization.\n\nCondition: Cannot edit own account.",
            "Action: Calls updateUserAccessAction(userId, values).\n• Enforces CHECK constraint: hospital roles must have facility_id; driver/admin must have NULL facility_id.\n• Re-evaluates receives_pickup_otp (requires phone_verified = true).\n• Writes security audit entry.",
            "Output: Green 'Saved' notification. User permissions take effect on their next API request or navigation. If deactivated, session terminates."
        ],
        [
            "[Provision User (Path B)]\n/admin/onboarding\nButton: Emerald Green",
            "Purpose: Directly creates and provisions an operational user without self-registration.\n\nCondition: Valid email & role.",
            "Action: Calls provisionUserAction(values).\n• Creates user record in database.\n• Triggers Clerk invitation email.\n• Pre-configures role and facility assignment.\n• Writes audit log.",
            "Output: User receives invitation email with secure activation link; appears in active Users table."
        ],
        [
            "[Approve Request]\nOnboarding Table\nButton: Emerald Green",
            "Purpose: Approves pending self-registration applicant.\n\nCondition: Status PENDING_REVIEW.",
            "Action: Calls approveOnboardingRequestAction(requestId).\n• Activates user account (active = true).\n• If applicant requested new hospital, automatically creates facility record.\n• Sets designated role.\n• Dispatches welcome SMS/Email via Outbox.",
            "Output: Request row disappears from pending list; applicant gains immediate operational access."
        ],
        [
            "[Confirm Reject]\nOnboarding Table\nButton: Danger Red",
            "Purpose: Denies onboarding request with written reason.\n\nCondition: Reason >= 3 characters.",
            "Action: Calls rejectOnboardingRequestAction(requestId, reason).\n• Updates status to REJECTED with rejection_reason.\n• Queues notification email explaining rejection.\n• Writes audit entry.",
            "Output: Applicant's pending screen updates to REJECTED. Account remains blocked from platform."
        ],
    ]
    add_button_table(doc, admin_user_buttons, "2.4 User Access & Onboarding Authorization Controls")

    # 2.5 Trip Planning & Dispatch
    admin_trip_buttons = [
        [
            "[Add Stop]\n/admin/trips/new\nButton: Emerald Ghost",
            "Purpose: Appends a waiting hospital to the current planned trip.\n\nCondition: Facility has waiting waste & ready = true.",
            "Action: Client state update appending facility to trip stops list.\n• Calculates cumulative container count and total payload weight (kg).\n• Checks against vehicle capacity limit and max_trip_stops setting.",
            "Output: Facility card moves to 'Stops in visiting order'. Total containers & weight recompute in real time."
        ],
        [
            "[Move Up / Down]\nStop Reorder Buttons\nIcons: ArrowUp / ArrowDown",
            "Purpose: Adjusts the sequential visiting order of collection stops.\n\nCondition: Stop list > 1.",
            "Action: Swaps array indices of adjacent stops in client state.\n• Updates stop order numbers (#1, #2, #3).",
            "Output: Visual reordering of route sequence. The resulting order determines driver arrival and manifest generation."
        ],
        [
            "[Remove Stop]\nStop List Row\nIcon: X (Rose)",
            "Purpose: Removes a facility from the planned trip.\n\nCondition: Stop in list.",
            "Action: Filters facility out of stops array.\n• Deducts containers and weight from trip totals.",
            "Output: Facility card returns to available Waiting List on the left panel."
        ],
        [
            "[Plan Trip]\n/admin/trips/new\nButton: Primary Blue",
            "Purpose: Commits route, locks stops, schedules waste, and dispatches trip to driver.\n\nCondition: Vehicle, driver, and >= 1 stop selected.",
            "Action: Calls createTripAction(payload).\n• Server re-verifies driver and vehicle licences (rejects if expired).\n• Enforces rule: driver and vehicle have zero other active trips.\n• Atomically sets trip PLANNED, stops PENDING, linked waste SCHEDULED.\n• Queues notification email to driver.",
            "Output: Redirects to /admin/trips/[id]. Trip appears immediately on Driver Mobile Console under 'Today's Trip'."
        ],
        [
            "[Reassign]\n/admin/trips/[id]\nButton: Secondary",
            "Purpose: Swaps vehicle or driver on a scheduled trip.\n\nCondition: Trip status PLANNED.",
            "Action: Calls /api/trips/[id]/reassign with idempotency key.\n• Cancels any active OTPs.\n• Re-validates licence of new driver/vehicle.\n• Revokes access of former driver immediately.\n• Writes audit log.",
            "Output: Trip details refresh. Former driver console clears; new driver receives notification."
        ],
        [
            "[Cancel Trip]\n/admin/trips/[id]\nButton: Danger Red",
            "Purpose: Aborts scheduled or uncollected trip with mandatory written rationale.\n\nCondition: Status PLANNED or IN_PROGRESS (no stop ARRIVED).",
            "Action: Calls /api/trips/[id]/cancel with reason.\n• Validates zero containers have been COLLECTED.\n• Sets trip CANCELLED.\n• Rolls all linked waste back from SCHEDULED to REGISTERED.\n• Cancels active pickup OTPs.\n• Writes audit log.",
            "Output: Trip status displays CANCELLED badge with cancel reason. Waste returns to Waiting List for future route planning."
        ],
        [
            "[Override Pickup]\n/admin/trips/[id]\nButton: Danger Red",
            "Purpose: Emergency administrative completion of stop when SMS and in-app confirmation fail.\n\nCondition: Stop ARRIVED, all containers scanned.",
            "Action: Calls /api/stops/[id]/override with written reason.\n• Validates EVERY container is SCANNED_AT_PICKUP.\n• Sets containers COLLECTED, stop COLLECTED_OVERRIDE, waste IN_TRANSIT.\n• Sets collection_method = 'OVERRIDE'.\n• Generates high-priority EFDA and Admin security alert.",
            "Output: Stop marked completed. Waste moves to IN_TRANSIT. Alert appears on regulator and admin security feeds."
        ],
        [
            "[Manifest (PDF)]\n/admin/trips/[id]\nLink: Secondary",
            "Purpose: Downloads official hazardous waste transport manifest for driver clipboard.\n\nCondition: Trip exists.",
            "Action: GET /api/trips/[id]/manifest.\n• Generates PDF listing vehicle plate, driver name, emergency contacts, stops, waste categories, and container count.",
            "Output: Browser opens or downloads formatted PDF Manifest Document ready for physical printing."
        ],
    ]
    add_button_table(doc, admin_trip_buttons, "2.5 Route Planning, Trip Dispatch & Override Controls")

    # 2.6 Approved Medicine Disposals Widget
    admin_disposal_widget = [
        [
            "[Assign Driver & Vehicle]\nApproved Disposals Widget\nButton: Amber",
            "Purpose: Assigns collection driver and fleet truck to EFDA-approved expired medicine batch.\n\nCondition: Status APPROVED or PENDING_DRIVER_ASSIGNMENT.",
            "Action: Calls assignDriverToRequestAction(reqId, driverId, vehicleId).\n• Links disposal request to driver profile.\n• Pre-selects driver's default vehicle.\n• Updates disposal request status to APPROVED.\n• Generates official Disposal QR Token.",
            "Output: Request status badge updates to 'Approved / Ready'. Request appears on assigned driver's console with route details."
        ],
        [
            "[View QR]\nDisposal Request Row\nButton: Secondary Ghost",
            "Purpose: Opens printable modal of the high-security cryptographic QR handover pass.\n\nCondition: Disposal request approved.",
            "Action: Client modal popup (DisposalQrModal).\n• Renders QR code embedding token, medicine lines, and facility code.",
            "Output: Modal displays QR code, container count, weight, and [Print QR Pass] button."
        ],
    ]
    add_button_table(doc, admin_disposal_widget, "2.6 Expired Medicine Disposal & Driver Assignment Controls")

    # 2.7 Financial Operations, Invoices & Contracts (ETB)
    admin_finance_buttons = [
        [
            "[Create Monthly Drafts]\n/admin/finance\nButton: Primary Blue",
            "Purpose: Manually executes monthly billing cycle for all active facility contracts.\n\nCondition: Super Admin role.",
            "Action: Calls runMonthlyInvoicesAction({ month }).\n• Scans active contracts.\n• Multiplies monthly_fee_santim by current VAT rate (e.g. 15%).\n• Creates unique DRAFT invoices (contract_id, period_start).\n• Skips contracts already billed for selected month.",
            "Output: Summary toast: 'XX drafts created, YY already existed'. Drafts appear in Invoices Ledger in ETB."
        ],
        [
            "[Issue]\nInvoice Table Row\nButton: Emerald Green",
            "Purpose: Formally issues draft invoice and establishes payment deadline.\n\nCondition: Invoice status DRAFT.",
            "Action: Calls issueInvoiceAction(invoiceId).\n• Sets status = ISSUED.\n• Computes due_date = issue_date + invoice_due_days.\n• Dispatches invoice notification email to Hospital Administrator.\n• Writes audit log.",
            "Output: StatusBadge changes to ISSUED with due date. Invoice total appears in hospital administrator billing dashboard."
        ],
        [
            "[Record Payment]\nInvoice Table Row\nButton: Sky Blue",
            "Purpose: Records offline bank wire / CBE / Telebirr settlement in ETB.\n\nCondition: Invoice status ISSUED.",
            "Action: Calls settleInvoiceAction(invoiceId, { reference, settledOn }).\n• Requires payment reference (bank slip #) and date.\n• Sets status = SETTLED_OFFLINE.\n• Clears overdue alert if invoice was delinquent.\n• Writes audit log.",
            "Output: Invoice displays SETTLED_OFFLINE badge with payment reference. Accounts receivable counters deduct settled amount."
        ],
        [
            "[Cancel Invoice]\nInvoice Table Row\nButton: Danger Red",
            "Purpose: Voids an incorrect or duplicate invoice with written reason.\n\nCondition: Status DRAFT or ISSUED.",
            "Action: Calls cancelInvoiceAction(invoiceId, { reason }).\n• Requires written reason.\n• Sets status = CANCELLED.\n• Writes audit log.",
            "Output: StatusBadge switches to CANCELLED. Amount removed from billing totals."
        ],
        [
            "[PDF]\nInvoice Table Row\nLink: Secondary",
            "Purpose: Generates official VAT billing invoice in Ethiopian Birr.\n\nCondition: Invoice exists.",
            "Action: GET /api/invoices/[id]/pdf.\n• Compiles PDF with Neba CS tax registration, facility TIN, period dates, monthly rate (ETB), VAT breakdown, and bank details.",
            "Output: Downloads printable tax-compliant invoice document."
        ],
        [
            "[New Contract]\n/admin/contracts\nButton: Primary Blue",
            "Purpose: Executes waste collection service agreement with health facility.\n\nCondition: Facility has no active contract.",
            "Action: Calls createContractAction(form).\n• Validates unique contract code (CT-YYYY-NNNN).\n• Converts monthly fee from ETB to santim.\n• Sets period dates and terms.\n• Status becomes ACTIVE.",
            "Output: Contract added to registry. Facility readiness flag clears 'No Contract' warning."
        ],
    ]
    add_button_table(doc, admin_finance_buttons, "2.7 Financial Operations, Invoices & Contracts (ETB)")

    # 2.8 Certificates & System Settings
    admin_system_buttons = [
        [
            "[Generate / Retry Now]\n/admin/certificates\nButton: Emerald Green",
            "Purpose: Manually triggers generation job for PENDING or FAILED certificates.\n\nCondition: >= 1 certificate waiting.",
            "Action: Calls runCertificateJobAction().\n• Compiles destruction certificate PDF.\n• Computes cryptographic SHA-256 hash.\n• Stores hash in certificates table.\n• Sets status = ISSUED.",
            "Output: Toast reports count of issued vs failed certificates. Public verify links activate immediately."
        ],
        [
            "[Revoke…]\nCertificate Table Row\nButton: Danger Red",
            "Purpose: Formally revokes an issued certificate due to discrepancy, with optional reissuance.\n\nCondition: Certificate status ISSUED.",
            "Action: Calls /api/certificates/[id]/revoke with reason and reissue toggle.\n• Marks certificate REVOKED.\n• If reissue = true, creates new PENDING certificate row for same destruction.\n• Note: Linked waste record status is NEVER rolled back.\n• Writes audit log.",
            "Output: StatusBadge switches to REVOKED. Public verify URL warns visitors that certificate is revoked."
        ],
        [
            "[Save]\nSettings Editor Row\nButton: Emerald Green",
            "Purpose: Updates global business setting (OTP expiry, cooldown, VAT rate, geofence radius).\n\nCondition: Value modified.",
            "Action: Calls updateSettingAction(key, value).\n• Validates data type (number, boolean, string list).\n• Updates settings table.\n• Future operations read new setting immediately (existing invoices preserve issued VAT rate).\n• Writes audit log.",
            "Output: Displays green 'Saved' indicator. New thresholds take effect platform-wide."
        ],
    ]
    add_button_table(doc, admin_system_buttons, "2.8 Certificates & Global System Settings Controls")

    # -------------------------------------------------------------
    # SECTION 3: HOSPITAL ADMINISTRATOR (/hospital)
    # -------------------------------------------------------------
    s3 = doc.add_heading(level=1)
    s3.paragraph_format.space_before = Pt(16)
    s3.paragraph_format.space_after = Pt(6)
    r_s3 = s3.add_run("3. Hospital Administrator Persona (/hospital) — Buttons & Workflows")
    r_s3.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "The Hospital Administrator manages biohazard waste compliance, container barcode manifests, staff permissions, "
        "and billing for their specific health facility (e.g. Dil Chora Referral Hospital). Every query is strictly filtered "
        "at the SQL layer by facility_id."
    )

    hosp_admin_buttons = [
        [
            "[Register Waste]\n/hospital/waste/new\nButton: Primary Blue",
            "Purpose: Creates a new categorized waste batch and generates container barcode labels.\n\nCondition: Category, count >= 1, weight > 0.",
            "Action: Calls registerWasteAction(values, idempotencyKey).\n• Validates input with Zod schema.\n• Inserts waste record with status REGISTERED.\n• Atomically creates N container records with status LABELLED.\n• Generates unique non-reusable container codes (WR-YYYY-NNNNN-CXX).\n• Writes audit log.",
            "Output: Form transitions to success view displaying container codes list and [Print Labels (PDF)] button."
        ],
        [
            "[Print Labels (PDF)]\nRegistration Success\nButton: Emerald Green",
            "Purpose: Generates and prints official adhesive barcode labels for waste bins.\n\nCondition: Waste registered.",
            "Action: GET /api/waste/[id]/labels.\n• Compiles PDF with container codes, Waste ID, facility name, MoH category, and high-density QR barcodes.\n• Includes Amharic text if label_amharic_enabled = true.",
            "Output: Opens printable PDF. Staff affix one barcode label to each physical container."
        ],
        [
            "[Confirm Handover in the App]\nTruck at Door Banner\nButton: Blue (Prominent)",
            "Purpose: In-app fallback confirmation when SMS pickup OTP fails to arrive at hospital.\n\nCondition: Stop ARRIVED, driver scanned ALL containers.",
            "Action: Calls /api/stops/[stopId]/confirm-handover.\n• Validates user belongs to same facility as stop.\n• Locks stop and re-checks that scannedContainers == expectedContainers.\n• Atomically sets containers COLLECTED, stop COLLECTED, waste IN_TRANSIT.\n• Sets collection_method = 'IN_APP'.\n• Writes audit log.",
            "Output: 'Truck at the door' banner clears with success checkmark. Waste status changes to 'In Transit'. Driver console completes stop."
        ],
        [
            "[Cancel Waste]\nWaste Table Row\nButton: Danger Rose",
            "Purpose: Cancels an uncollected waste record with mandatory written reason.\n\nCondition: Status REGISTERED or SCHEDULED (not ARRIVED).",
            "Action: Calls cancelWasteAction(id, { reason }).\n• Refuses with HTTP 409 if truck is currently ARRIVED at facility.\n• Updates waste status to CANCELLED.\n• If part of PLANNED trip stop, cancels stop waste link.\n• Writes audit log.",
            "Output: Waste row status updates to CANCELLED. Cannot be scheduled on future collection trips."
        ],
        [
            "[Toggle Pickup Codes]\n/hospital/users\nCheckbox: User Row",
            "Purpose: Authorizes or revokes whether a specific staff member receives SMS pickup OTPs.\n\nCondition: Staff phone verified in Clerk.",
            "Action: Calls updateHospitalUserAction(userId, { receivesPickupOtp }).\n• Enforces business rule: user must have phone_verified = true.\n• Updates receives_pickup_otp column in database.\n• Writes audit log.",
            "Output: Checkbox updates. When driver marks arrival, system sends SMS OTP to all active verified recipients."
        ],
        [
            "[Toggle Active Status]\n/hospital/users\nCheckbox: User Row",
            "Purpose: Instantly suspends or restores facility staff login access.\n\nCondition: Cannot disable own account.",
            "Action: Calls updateHospitalUserAction(userId, { active }).\n• Updates active column in users table.\n• If inactive, user is immediately blocked and redirected to /pending.",
            "Output: Checkbox updates. Inactive staff cannot log in or confirm handovers."
        ],
        [
            "[Approve Request]\nDisposal Requests List\nButton: Emerald Green",
            "Purpose: Approves hospital pharmacy expired medicine disposal request.\n\nCondition: Status PENDING_APPROVAL.",
            "Action: Calls decideRequestAction(reqId, { decision: 'APPROVE' }).\n• If batch contains expired medicine, routes to PENDING_EFDA_APPROVAL.\n• If no regulatory review needed, sets status to PENDING_DRIVER_ASSIGNMENT.\n• Writes audit log.",
            "Output: Request status badge updates. Notification sent to EFDA officer or Super Admin."
        ],
        [
            "[Reject…]\nDisposal Requests List\nButton: Danger Rose",
            "Purpose: Rejects medicine disposal request with written explanation.\n\nCondition: Reason >= 3 characters.",
            "Action: Calls decideRequestAction(reqId, { decision: 'REJECT', reason }).\n• Sets request status to REJECTED with decision_reason.\n• Restores medicine stock line balances.\n• Writes audit log.",
            "Output: Request marked REJECTED. Submitting staff sees rejection rationale."
        ],
        [
            "[Download Certificate]\n/hospital/certificates\nLink: Emerald",
            "Purpose: Retrieves official PDF Certificate of Destruction for audited compliance.\n\nCondition: Certificate status ISSUED.",
            "Action: GET /api/certificates/[id]/pdf.\n• Streams verified PDF signed with SHA-256 cryptographic stamp.",
            "Output: Downloads official government-recognized Destruction Certificate."
        ],
    ]
    add_button_table(doc, hosp_admin_buttons, "Table 3.1: Hospital Administrator Interactive Controls")

    # -------------------------------------------------------------
    # SECTION 4: HOSPITAL STAFF / EHO (/hospital)
    # -------------------------------------------------------------
    s4 = doc.add_heading(level=1)
    s4.paragraph_format.space_before = Pt(16)
    s4.paragraph_format.space_after = Pt(6)
    r_s4 = s4.add_run("4. Hospital Staff / EHO Persona (/hospital) — Buttons & Workflows")
    r_s4.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "Hospital Staff and Environmental Health Officers (EHO) handle day-to-day biohazard waste packaging, "
        "medicine stock logging, disposal request formulation, and physical handover to collection drivers."
    )

    hosp_staff_buttons = [
        [
            "[Register Waste]\n/hospital/waste/new\nButton: Primary Blue",
            "Purpose: Registers daily biohazard waste bags/boxes by MoH color code.\n\nCondition: Valid category, count, weight.",
            "Action: Form submission calling registerWasteAction().\n• Creates waste record and container barcode rows.\n• Audited in database.",
            "Output: Confirmation screen with container codes list and barcode printing trigger."
        ],
        [
            "[Check File]\n/hospital/stock/upload\nButton: Secondary",
            "Purpose: Previews and validates medicine inventory spreadsheet before importing.\n\nCondition: .xlsx or .csv selected.",
            "Action: Calls previewStockAction(formData).\n• Validates required column headers.\n• Validates expiry date formats and positive quantities.\n• Flags duplicate batches or syntax errors.",
            "Output: Displays preview table with valid rows vs error rows. Button [Save Lines] appears if zero errors exist."
        ],
        [
            "[Save Lines]\n/hospital/stock/upload\nButton: Emerald Green",
            "Purpose: Commits validated medicine stock inventory to hospital database.\n\nCondition: File validated with zero errors.",
            "Action: Calls saveStockAction({ lines }).\n• Upserts stock lines by batch number.\n• Updates quantities and expiration dates.\n• Writes audit log.",
            "Output: Success toast: 'Saved: XX new lines, YY updated'. Stock lines appear in inventory table with expiry warnings."
        ],
        [
            "[Submit Disposal Request]\n/hospital/stock/requests/new\nButton: Primary Blue",
            "Purpose: Formulates formal disposal batch of expired or damaged pharmaceuticals.\n\nCondition: >= 1 stock item checked.",
            "Action: Calls createDisposalRequestAction(payload, idempotencyKey).\n• Deducts requested quantities from active stock.\n• Sets request status to PENDING_APPROVAL.\n• Routes request to Hospital Administrator.\n• Writes audit log.",
            "Output: Toast: 'Request MD-YYYY-NNNN sent for approval'. Appears on hospital admin and EFDA review queues."
        ],
        [
            "[View QR Pass]\nDisposal Requests Row\nButton: Secondary Ghost",
            "Purpose: Opens printable QR handover pass for packaging on disposal bins.\n\nCondition: Request approved.",
            "Action: Opens DisposalQrModal component.\n• Renders QR code embedding token, medicine lines, and facility code.",
            "Output: Displays high-density QR code for driver scan at pickup."
        ],
    ]
    add_button_table(doc, hosp_staff_buttons, "Table 4.1: Hospital Staff & EHO Operational Controls")

    # -------------------------------------------------------------
    # SECTION 5: FLEET DRIVER (/driver)
    # -------------------------------------------------------------
    s5 = doc.add_heading(level=1)
    s5.paragraph_format.space_before = Pt(16)
    s5.paragraph_format.space_after = Pt(6)
    r_s5 = s5.add_run("5. Fleet Driver Persona (/driver) — Buttons & Workflows")
    r_s5.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "The Fleet Driver console is a mobile-first, high-contrast interface designed for touch interaction in truck cabs. "
        "It guides the driver through sequential stops, automated GPS geofencing, QR barcode scanning, and cryptographic OTP verification."
    )

    driver_buttons = [
        [
            "[Start Trip TR-YYYY-NNNN]\nDriver Console Header\nButton: Amber (Big Touch)",
            "Purpose: Initiates scheduled daily collection route.\n\nCondition: Driver has trip assigned with status PLANNED.",
            "Action: Calls /api/trips/[id]/start with idempotency key.\n• Server re-verifies driver licence and vehicle licence.\n• Ensures vehicle has zero other active trips.\n• Atomically updates trip status to IN_PROGRESS.\n• Writes audit log.",
            "Output: Trip banner turns active with live stop sequence. Map shows route. Stop #1 unlocks with 'I have arrived' button."
        ],
        [
            "[I have arrived]\nStop Card (#1)\nButton: Amber (Big Touch)",
            "Purpose: Signals physical arrival at hospital loading dock and triggers SMS OTP.\n\nCondition: Stop status PENDING, previous stop completed.",
            "Action: Calls /api/stops/[stopId]/arrive.\n• Captures latest vehicle GPS coordinate.\n• Compares distance against facility geofence (geofence_radius_m).\n• Stores arrival_distance_m and geofence_ok (flag only, non-blocking).\n• Sets stop status = ARRIVED.\n• Generates 6-digit cryptographic OTP.\n• Queues SMS via Outbox to all verified hospital OTP recipients.",
            "Output: Stop card expands. Status changes to ARRIVED. Displays SMS delivery countdown timer. Container scanner activates."
        ],
        [
            "[Scan Container QR]\nScanner Component\nButton: Camera Viewfinder",
            "Purpose: Uses smartphone camera to scan unique barcode on each physical waste container.\n\nCondition: Stop status ARRIVED.",
            "Action: BarcodeDetector detects QR containing container code (WR-YYYY-NNNNN-CXX).\n• Submits POST /api/stops/[stopId]/scan with code.\n• Validates container belongs to this stop and is LABELLED.\n• Sets container status = SCANNED_AT_PICKUP.\n• Inserts handover row.\n• Device vibrates (80ms haptic confirmation).",
            "Output: Counter increments: 'Scanned X of N'. Container code listed with green checkmark. (Does NOT complete pickup)."
        ],
        [
            "[Manual Code Fallback]\nScanner Component\nInput & Submit: [Scan]",
            "Purpose: Manual container code entry if camera lens is damaged or barcode torn.\n\nCondition: Stop status ARRIVED.",
            "Action: Submits typed container code through identical validation pipeline as camera scan.",
            "Output: Increments scanned count if code matches expected manifest; displays error toast if code is invalid."
        ],
        [
            "[Confirm (Pickup Code)]\nOTP Entry Form\nButton: Emerald Green",
            "Purpose: Submits 6-digit SMS OTP provided by hospital staff to complete pickup.\n\nCondition: ALL containers scanned (scanned == expected), 6 digits entered.",
            "Action: Calls /api/stops/[stopId]/verify-otp with code.\n• Compares code against stored SHA-256 hash in constant time.\n• Enforces max attempts (otp_max_attempts).\n• Atomically sets: all containers = COLLECTED, stop = COLLECTED, waste = IN_TRANSIT.\n• Sets collection_method = 'OTP'.\n• Writes audit log.",
            "Output: Stop card collapses with green 'Collected' badge. Next stop on route unlocks. Waste moves to IN_TRANSIT on all dashboards."
        ],
        [
            "[Send a new code]\nStop Actions\nButton: White Outline",
            "Purpose: Resends SMS OTP to hospital staff if original message was delayed.\n\nCondition: Subject to 60s cooldown, max 3 codes per stop.",
            "Action: Calls /api/stops/[stopId]/resend-otp.\n• Cancels previous OTP hash.\n• Generates fresh 6-digit OTP.\n• Queues SMS via Outbox.",
            "Output: Info message: 'A new code was sent to the hospital.' Expiration countdown timer resets to 10 minutes."
        ],
        [
            "[Skip this stop…]\nStop Actions\nButton: Rose Outline",
            "Purpose: Aborts collection at facility due to obstruction, closure, or unready waste.\n\nCondition: Stop PENDING or ARRIVED.",
            "Action: Expands reason dialogue requiring >= 3 characters.\n• Driver clicks [Skip stop].\n• Calls /api/stops/[stopId]/skip with written reason.\n• Sets stop status = SKIPPED with skip_reason.\n• Any scanned containers roll back to LABELLED.\n• Linked waste rolls back to REGISTERED.\n• Cancels active OTPs.\n• Writes audit log.",
            "Output: Stop card collapses with SKIPPED status. Waste returns to Waiting List for future trip planning. Next stop activates."
        ],
        [
            "[Scan QR Code for Waste Pickup]\nAssigned Disposals Panel\nButton: Amber",
            "Purpose: Direct QR scan for approved medicine disposal package handover.\n\nCondition: Disposal request assigned to driver.",
            "Action: Calls scanRequestByDriverAction(qrCode).\n• Validates QR token matches assigned request.\n• Updates status to RECEIVED_BY_DRIVER.\n• Records driver_scanned_at timestamp.\n• Writes audit log.",
            "Output: Toast: 'QR Code verified! Request marked as Received by Driver.' Card shows green scan confirmation."
        ],
    ]
    add_button_table(doc, driver_buttons, "Table 5.1: Fleet Driver Mobile Console Interactive Controls")

    make_callout(
        doc,
        "Zero Partial Pickup Rule",
        "The SafeWaste engine strictly forbids partial pickups. If a facility registers 5 containers, the [Confirm] button "
        "and in-app confirmation remain permanently disabled until exactly 5 containers are scanned. If even one container "
        "is missing or damaged at the hospital, the driver MUST click [Skip this stop…]. The entire batch rolls back safely.",
        "warning"
    )

    # -------------------------------------------------------------
    # SECTION 6: DISPOSAL OPERATOR (/disposal)
    # -------------------------------------------------------------
    s6 = doc.add_heading(level=1)
    s6.paragraph_format.space_before = Pt(16)
    s6.paragraph_format.space_after = Pt(6)
    r_s6 = s6.add_run("6. Disposal Facility Operator Persona (/disposal) — Buttons & Workflows")
    r_s6.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "The Disposal Facility Operator oversees the central treatment plant. Responsibilities include checking in inbound containers "
        "from collection trucks, logging discrepancies (missing/refused containers), executing irreversible destruction records, "
        "and generating official regulatory certificates."
    )

    disposal_buttons = [
        [
            "[Check in Container (Scan)]\nReceipt Desk Component\nCamera / Manual Input",
            "Purpose: Scans each container offloaded from collection truck into facility yard.\n\nCondition: Container status COLLECTED.",
            "Action: Calls /api/disposal/receive with containerCode.\n• Validates container is COLLECTED.\n• Updates container status to RECEIVED.\n• Checks remaining containers on truck.\n• When all containers for waste record are checked in, waste transitions from IN_TRANSIT to RECEIVED.\n• When all containers on truck are checked in, trip transitions to COMPLETED.\n• Writes audit log.",
            "Output: Success toast: 'Received · WR-YYYY-NNNN: X still on truck · trip completed'. Container moves to Destruction Queue."
        ],
        [
            "[Missing / Refused…]\nOn-Board Container Row\nButton: Secondary Ghost",
            "Purpose: Flags an expected container that failed physical inspection or went missing.\n\nCondition: Container on truck manifest.",
            "Action: Expands ExceptionForm.\n• Operator selects 'Missing (not on truck)' or 'Refused (damaged/non-compliant)'.\n• Inputs mandatory written explanation (>= 3 chars).\n• Calls /api/disposal/exception.\n• Sets container status = MISSING or REFUSED.\n• Triggers immediate High-Priority Alert to Super Admin and EFDA.\n• Writes audit log.",
            "Output: Container removed from inbound list. Exception logged. Discrepancy counter increments on regulator dashboard."
        ],
        [
            "[Record Destruction]\nDestruction Desk Row\nButton: Emerald Green",
            "Purpose: Enters destruction parameters and initiates irreversible disposal confirmation.\n\nCondition: Waste status RECEIVED, >= 1 container received.",
            "Action: Operator selects treatment method (Incineration, Autoclaving, Chemical Treatment, Encapsulation).\n• Clicks [Record destruction].",
            "Output: Replaces button with prominent confirmation banner: 'Destroy X containers? This cannot be undone.' and [Yes, destroyed] button."
        ],
        [
            "[Yes, Destroyed]\nDestruction Confirmation\nButton: Danger Rose",
            "Purpose: Irreversibly locks destruction record and enqueues official certificate generation.\n\nCondition: Confirmation banner active.",
            "Action: Calls /api/disposal/waste/[id]/destroy with methodCode.\n• Atomically marks all received containers DESTROYED.\n• Sets waste status = DISPOSED (or DISPOSED_WITH_EXCEPTION if missing containers exist).\n• Permanently locks disposal record.\n• Enqueues asynchronous PDF Certificate generation job.\n• Writes audit log.",
            "Output: Card displays success: 'Destroyed. Certificate CT-YYYY-NNNN is being prepared.' Waste moves to archived ledger."
        ],
        [
            "[Close Without Destruction]\nDestruction Desk Row\nButton: Danger Rose",
            "Purpose: Closes a waste record where ALL containers were flagged MISSING or REFUSED.\n\nCondition: Received count == 0.",
            "Action: Calls /api/disposal/waste/[id]/close with mandatory written reason.\n• Sets waste status = CLOSED_NO_DESTRUCTION.\n• Generates critical security alert.\n• NO certificate is generated (audit record only).\n• Writes audit log.",
            "Output: Record closed without certificate. System alerts EFDA and Super Admin of total destruction failure."
        ],
        [
            "[Scan QR & Receive at Facility]\nDisposal QR Scanner\nButton: Primary Blue",
            "Purpose: Receives approved medicine disposal batch at central facility via QR pass.\n\nCondition: Status RECEIVED_BY_DRIVER.",
            "Action: Calls scanRequestByOperatorAction(qrCode).\n• Validates driver scan milestone is complete.\n• Sets status = RECEIVED_AT_FACILITY.\n• Records operator_scanned_at timestamp.\n• Writes audit log.",
            "Output: Dual-scan verification trail marks Operator Scan as complete with green checkmark. [Mark as Disposed] unlocks."
        ],
        [
            "[Mark as Disposed & Generate PDF]\nDisposal QR Scanner\nButton: Emerald Gradient",
            "Purpose: Finalizes incineration of medicine batch and creates public certificate.\n\nCondition: Status RECEIVED_AT_FACILITY.",
            "Action: Calls markDisposalCompleteAction(requestId).\n• Sets status = DISPOSED.\n• Generates unique certificate number.\n• Compiles destruction certificate PDF.\n• Computes SHA-256 hash.\n• Writes audit log.",
            "Output: Status updates to DISPOSED. Green banner displays certificate number and active [Download PDF Certificate] button."
        ],
    ]
    add_button_table(doc, disposal_buttons, "Table 6.1: Disposal Facility Operator Interactive Controls")

    # -------------------------------------------------------------
    # SECTION 7: EFDA REGULATORY OFFICER (/regulator)
    # -------------------------------------------------------------
    s7 = doc.add_heading(level=1)
    s7.paragraph_format.space_before = Pt(16)
    s7.paragraph_format.space_after = Pt(6)
    r_s7 = s7.add_run("7. EFDA Regulatory & Audit Officer (/regulator) — Buttons & Workflows")
    r_s7.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "The EFDA Regulatory & Audit Officer provides federal compliance oversight across all 18 regional health facilities. "
        "In strict compliance with platform rules, this role is READ-ONLY across all operational data surfaces, with exactly "
        "ONE write permission: approving or rejecting medicine disposal authorization requests."
    )

    regulator_buttons = [
        [
            "[Verify the Hash Chain]\n/regulator/audit\nButton: Emerald Green",
            "Purpose: Recomputes and cryptographically verifies the sequential SHA-256 integrity of the entire audit ledger.\n\nCondition: Regulator role.",
            "Action: Calls checkAuditChainAction().\n• Iterates through audit_logs table ordered by sequence number.\n• Sequentially recalculates SHA-256(prev_hash + payload + timestamp).\n• Validates hash matches stored record.\n• Detects any SQL tampering, deletion, or record modification.",
            "Output: If intact: green shield icon with 'Chain intact: XX entries checked.' If tampered: red alert displaying exact sequence number where hash broke."
        ],
        [
            "[Approve Request]\n/regulator/approvals\nButton: Emerald Green",
            "Purpose: Grants formal federal regulatory authorization to destroy expired/controlled pharmaceuticals.\n\nCondition: Status PENDING_EFDA_APPROVAL.",
            "Action: Calls decideRequestAction(reqId, { decision: 'APPROVE' }).\n• Updates status to PENDING_DRIVER_ASSIGNMENT (or APPROVED).\n• Generates official EFDA disposal permit.\n• Routes request to Super Admin for driver assignment.\n• Writes audit log.",
            "Output: Request status updates to 'Approved'. Request is unlocked for collection and transport to incinerator."
        ],
        [
            "[Reject…]\n/regulator/approvals\nButton: Danger Rose",
            "Purpose: Denies disposal authorization with written regulatory explanation.\n\nCondition: Reason >= 3 characters.",
            "Action: Calls decideRequestAction(reqId, { decision: 'REJECT', reason }).\n• Updates status to REJECTED with decision_reason.\n• Restores hospital inventory line balances.\n• Dispatches regulatory refusal notification to hospital administrator.\n• Writes audit log.",
            "Output: Request marked REJECTED. Hospital cannot destroy batch without re-applying with corrected documentation."
        ],
        [
            "[Export CSV]\nAudit / League Tables\nLink: Secondary Outline",
            "Purpose: Exports regulatory compliance records, waste volumes, or audit ledger to CSV spreadsheet.\n\nCondition: Data present.",
            "Action: Client-side CSV generator compiling active table view.\n• Formats headers, dates, ETB totals, and verification tokens.\n• Triggers browser file download.",
            "Output: CSV spreadsheet file downloaded for regional health bureau archiving."
        ],
        [
            "[Play / Pause Replay]\n/regulator/map\nMedia Controls: Player",
            "Purpose: Replays historical GPS breadcrumb tracking of completed collection trips.\n\nCondition: Historical trip selected.",
            "Action: Advances map playback timeline.\n• Animates vehicle position along actual recorded tracker coordinates.\n• Displays vehicle speed and timestamps at each hospital stop.",
            "Output: Visual route replay verifying physical truck movement and stops."
        ],
    ]
    add_button_table(doc, regulator_buttons, "Table 7.1: EFDA Regulatory & Compliance Controls")

    # -------------------------------------------------------------
    # SECTION 8: PUBLIC VERIFICATION & ONBOARDING
    # -------------------------------------------------------------
    s8 = doc.add_heading(level=1)
    s8.paragraph_format.space_before = Pt(16)
    s8.paragraph_format.space_after = Pt(6)
    r_s8 = s8.add_run("8. Public Verification & Self-Registration Portals — Buttons & Workflows")
    r_s8.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "Public endpoints require no authentication. They allow citizens, municipal inspectors, and healthcare auditors to verify "
        "destruction certificates and allow new staff to apply for platform onboarding."
    )

    public_buttons = [
        [
            "[Verify]\nLanding Page (Home)\nButton: Emerald Green",
            "Purpose: Quick lookup of certificate from pasted QR link or token.\n\nCondition: Valid token string.",
            "Action: Parses token from URL or text string.\n• Validates token regex (^[A-Za-z0-9_-]{22,100}$).\n• Navigates browser to /verify/[token].",
            "Output: Redirects user directly to public certificate verification card."
        ],
        [
            "[Download Certificate (PDF)]\n/verify/[token]\nButton: Emerald Green",
            "Purpose: Downloads official Destruction Certificate for public archiving.\n\nCondition: Certificate status ISSUED (hidden if REVOKED).",
            "Action: GET /api/verify/[token]/pdf.\n• Streams PDF certificate containing official seal, container metrics, and SHA-256 hash.",
            "Output: Downloads PDF file."
        ],
        [
            "[Compare File]\nFileCheck Component\nButton: Emerald Green",
            "Purpose: Uploads a physical PDF certificate to verify whether it matches the official digital record.\n\nCondition: .pdf file selected.",
            "Action: POST /api/verify/[token]/check-file (multipart/form-data).\n• Computes SHA-256 hash of uploaded file in memory (file is immediately discarded).\n• Compares computed hash with stored hash in database.",
            "Output: If identical: green banner 'This file is the original certificate CT-YYYY-NNNN.' If altered: red banner 'This file does NOT match the certificate on record. It may have been modified.'"
        ],
        [
            "[Submit Application]\n/pending (Self-Registration)\nButton: Primary Blue",
            "Purpose: Submits user self-registration application for administrative approval.\n\nCondition: Name, email, role, and facility selected.",
            "Action: Calls submitOnboardingRequestAction(payload).\n• Creates record in onboarding_requests table with status PENDING_REVIEW.\n• Profile is assigned ZERO operational access.\n• Dispatches alert to Super Admin.",
            "Output: Screen switches to 'Under Review' status card with clock icon. User remains blocked until approved."
        ],
    ]
    add_button_table(doc, public_buttons, "Table 8.1: Public Verification & Self-Registration Controls")

    # -------------------------------------------------------------
    # SECTION 9: END-TO-END WORKFLOW LIFECYCLE
    # -------------------------------------------------------------
    s9 = doc.add_heading(level=1)
    s9.paragraph_format.space_before = Pt(16)
    s9.paragraph_format.space_after = Pt(6)
    r_s9 = s9.add_run("9. Complete End-to-End Operational Lifecycle Trace")
    r_s9.font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph(
        "The following structured sequence traces a batch of hazardous healthcare waste from the moment of hospital "
        "registration, through collection, transit, and incineration, to public verification:"
    )

    lifecycle_steps = [
        ("Step 1: Hospital Registration", "Hospital Staff clicks [Register Waste] at /hospital/waste/new. Enters 4 containers of Infectious Waste weighing 48.5 kg. Database creates waste row REGISTERED and 4 container rows LABELLED. Staff clicks [Print Labels (PDF)] and affixes barcodes."),
        ("Step 2: Route Dispatch", "Super Admin clicks [Add Stop] in /admin/trips/new. Adds hospital to route with Driver Abebe and Truck Plate 3-54321. Admin clicks [Plan Trip]. Trip becomes PLANNED; waste becomes SCHEDULED. Manifest PDF generated."),
        ("Step 3: Route Start", "Driver Abebe opens /driver console on mobile. Clicks [Start Trip]. Trip becomes IN_PROGRESS. Hardware GPS tracker begins streaming coordinates to /api/tracking/live. Fleet radar displays live truck marker."),
        ("Step 4: Hospital Arrival", "Driver arrives at hospital loading dock and clicks [I have arrived]. System verifies GPS coordinates are within 100m geofence. Stop status becomes ARRIVED. Outbox immediately dispatches SMS OTP code to verified hospital staff."),
        ("Step 5: Handover & QR Scanning", "Driver clicks [Scan Container QR] and scans all 4 containers with phone camera. Each container becomes SCANNED_AT_PICKUP. Counter reaches '4 of 4'. Hospital staff reads 6-digit OTP code from phone. Driver enters code and clicks [Confirm]. Atomically, containers become COLLECTED, stop becomes COLLECTED, waste becomes IN_TRANSIT. (If SMS delayed, hospital admin clicks [Confirm Handover in the App])."),
        ("Step 6: Disposal Site Reception", "Truck arrives at Central Treatment Plant. Disposal Operator opens /disposal terminal and clicks [Check in Container]. Scans all 4 containers. Containers become RECEIVED. Waste becomes RECEIVED. All truck containers cleared; trip automatically marked COMPLETED."),
        ("Step 7: Destruction & Incineration", "Disposal Operator clicks [Record Destruction] on the received batch. Selects 'High-Temperature Incineration (1200°C)'. Clicks [Yes, Destroyed]. Containers become DESTROYED; waste becomes DISPOSED. Disposal record permanently locked. System queues PDF certificate generation job."),
        ("Step 8: Certificate Issuance & Public Verification", "Certificate job generates PDF with cryptographic SHA-256 hash. Status becomes ISSUED. Hospital and EFDA dashboards update with [Download Certificate] button. Public visitor scans QR code on certificate, opening /verify/[token], and uploads file to [Compare File] to verify tamper-proof authenticity."),
        ("Step 9: Monthly Billing & Audit", "At end of month, Super Admin clicks [Create Monthly Drafts]. Generates invoice for hospital contract in Ethiopian Birr (ETB) with 15% VAT. Super Admin clicks [Issue]. Later, upon bank wire receipt, clicks [Record Payment] with CBE slip reference. EFDA Officer clicks [Verify the Hash Chain] in /regulator/audit to confirm ledger remains 100% intact.")
    ]

    for title, desc in lifecycle_steps:
        p_step = doc.add_paragraph()
        p_step.paragraph_format.space_before = Pt(4)
        p_step.paragraph_format.space_after = Pt(2)
        p_step.paragraph_format.line_spacing = 1.15
        r_st = p_step.add_run(f"{title}: ")
        r_st.bold = True
        r_st.font.name = "Calibri"
        r_st.font.size = Pt(10)
        r_st.font.color.rgb = RGBColor(15, 118, 110)
        r_sd = p_step.add_run(desc)
        r_sd.font.name = "Calibri"
        r_sd.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    make_callout(
        doc,
        "Conclusion & Operational Governance",
        "Every button in SafeWaste is built for zero data loss, strict role scoping, and complete regulatory transparency. "
        "For system support, licence renewals, or configuration updates, contact the SafeWaste platform administrative team.",
        "success"
    )

    # Save document
    output_path = os.path.join(os.getcwd(), "docs", "SafeWaste-Complete-Button-and-Workflow-Manual.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    build_document()
